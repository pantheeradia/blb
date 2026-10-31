from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
OUT_PATH = OUT_DIR / "Bean_Leaf_Brew_AEO_SEO_Audit_and_Implementation_Framework.docx"


BLACK = "000000"
COFFEE = "4A2512"
MID_BROWN = "8B5E3C"
PALE_CREAM = "F7F0E6"
PALE_BLUE = "EDF3F7"
PALE_GRAY = "F3F4F6"
WHITE = "FFFFFF"
GRID = "D9D9D9"
MUTED = "5B6573"


def set_run_font(run, name: str = "Aptos", size: float | None = None, color: str | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=105, start=115, bottom=105, end=115):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, width_inches: float):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_hyperlink(paragraph, text: str, url: str, color=MID_BROWN):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Aptos")
    r_fonts.set(qn("w:hAnsi"), "Aptos")
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=8.5, color=MUTED)


def add_para(doc: Document, text: str = "", style: str | None = None, bold_lead: str | None = None, keep=False):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    if keep:
        keep_with_next(p)
    return p


def add_bullets(doc: Document, items: Iterable[str], level: int = 0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    for item in items:
        p = doc.add_paragraph(style=style)
        r = p.add_run(item)
        set_run_font(r)


def add_numbered(doc: Document, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        set_run_font(r)


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[float] | None = None, font_size: float = 8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, (cell, header) in enumerate(zip(hdr.cells, headers)):
        set_cell_shading(cell, COFFEE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            set_cell_width(cell, widths[idx])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=font_size, color=WHITE, bold=True)
    for row_idx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for idx, (cell, value) in enumerate(zip(row.cells, row_data)):
            set_cell_shading(cell, WHITE if row_idx % 2 == 0 else PALE_CREAM)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_cell_width(cell, widths[idx])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if idx == 0 and len(headers) <= 4:
                r = p.add_run(str(value))
                set_run_font(r, size=font_size, bold=True)
            else:
                r = p.add_run(str(value))
                set_run_font(r, size=font_size)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    return table


def add_small_label(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text.upper())
    set_run_font(r, size=8.5, color=MID_BROWN, bold=True)
    return p


def page_break(doc: Document):
    doc.add_page_break()


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(9.7)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(10)

    for style_name, size, before, after in (
        ("Heading 1", 19, 16, 7),
        ("Heading 2", 13.5, 12, 5),
        ("Heading 3", 10.5, 9, 4),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(9.4)
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.08


def configure_sections(doc: Document):
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)
        section.header_distance = Inches(0.28)
        section.footer_distance = Inches(0.28)
        footer = section.footer
        p = footer.paragraphs[0]
        p.clear()
        left = p.add_run("Bean Leaf Brew AEO and SEO Framework    ")
        set_run_font(left, size=8.2, color=MUTED)
        add_page_number(p)


def add_section_intro(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    r = p.add_run(text)
    set_run_font(r, size=10.1, color=MUTED)
    return p


def build_document() -> Document:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    # Cover page
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("BEAN LEAF BREW")
    set_run_font(r, size=9.5, color=MID_BROWN, bold=True)

    title = doc.add_paragraph(style="Title")
    title.add_run("AEO and SEO Audit and Implementation Framework")

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(20)
    r = sub.add_run("Comprehensive practices for search engines answer engines and AI discovery")
    set_run_font(r, name="Aptos Display", size=15, color=BLACK)

    add_para(
        doc,
        "This handbook defines the complete method for assessing Bean Leaf Brew's current organic and AI-search visibility, prioritizing the work, and implementing improvements. It covers conventional search, local and shopping surfaces, Google generative search, Bing and Copilot, ChatGPT Search, Perplexity, Claude search, and emerging agentic discovery.",
    )
    add_para(
        doc,
        "The main conclusion is practical: AEO is strongest when it extends sound SEO. The site must first be accessible, indexable, understandable, trustworthy, and useful. AEO then adds explicit source quality, answer coverage, citation measurement, crawler policy, entity consistency, and cross-platform testing. No tactic guarantees a ranking or citation.",
    )

    add_small_label(doc, "How to use this handbook")
    add_table(
        doc,
        ["Stage", "Purpose", "Primary output"],
        [
            ["1 Audit", "Measure SEO and AEO health using technical evidence, content review, platform data, prompt testing, and competitor comparison.", "Baseline score, issue register, page inventory, prompt benchmark, and prioritized backlog."],
            ["2 Implementation", "Fix blockers first, improve high-value pages, build answer-ready content and entity signals, then operate a repeatable measurement program.", "Technical releases, content upgrades, structured data, feeds, authority work, dashboards, and governance cadence."],
        ],
        widths=[1.2, 3.2, 2.45],
        font_size=8.5,
    )

    add_small_label(doc, "Document status")
    add_para(doc, "Prepared 30 October 2026 for planning and execution. Platform-specific controls and reports should be rechecked before each major release because AI search products change quickly.")

    page_break(doc)

    add_heading(doc, "Executive Summary", 1)
    add_section_intro(doc, "The framework joins SEO and AEO into one operating system while keeping their different outcomes measurable.")

    add_heading(doc, "What success means", 2)
    add_para(doc, "SEO success means qualified pages are discovered, indexed, ranked, clicked, and converted. AEO success means the brand and its pages are retrieved, cited or used accurately in generated answers, and then contribute to qualified visits, consideration, assisted conversions, or brand demand. A citation is useful evidence, but it is not the business outcome by itself.")

    add_heading(doc, "Core principles", 2)
    add_bullets(doc, [
        "Fix eligibility before optimization. A blocked, non-indexable, duplicated, thin, or poorly rendered page is not a reliable candidate for either search result visibility or AI citation.",
        "Publish information worth citing. Original observations, clear product facts, expert guidance, first-hand tests, transparent methods, current policies, useful images, and primary data are stronger than rewritten commodity content.",
        "Make the source easy to understand. Use descriptive titles and headings, concise answer passages, clear tables where comparison is needed, consistent entity names, crawlable links, and structured data that matches visible content.",
        "Separate access decisions from training decisions. Major AI providers use different user agents for search, model training, and user-initiated retrieval. The policy should be explicit and reviewed with legal and security owners.",
        "Measure a pipeline, not one score. Track crawling, indexation, search visibility, AI feature visibility, citations, brand accuracy, referrals, assisted revenue, and conversion quality.",
        "Treat AEO tests as probabilistic. AI answers vary by engine, wording, location, personalization, and time. Repeat tests and retain prompts, dates, outputs, citations, and screenshots or exports.",
    ])

    add_heading(doc, "Recommended first outcome", 2)
    add_para(doc, "Complete a baseline audit before changing the site. The audit should identify critical blockers, establish a page and prompt benchmark, and produce a 90-day backlog. Implementation should begin with measurement, crawl and index controls, canonical product URLs, page metadata, product data accuracy, Merchant Center and Business Profile hygiene, Organization and Product markup, and upgrades to the pages most likely to generate revenue or citations.")

    add_heading(doc, "Document map", 2)
    add_table(
        doc,
        ["Part", "Contents"],
        [
            ["One", "Definitions, evidence model, audit scope, scorecard, and detailed SEO and AEO audit practices."],
            ["Two", "Implementation principles, 90-day roadmap, workstreams, ownership, and operating cadence."],
            ["Three", "Page patterns, crawler policy, measurement dictionary, myths, checklists, and research sources."],
        ],
        widths=[1.0, 5.85],
        font_size=8.7,
    )

    page_break(doc)
    add_heading(doc, "Part One Audit Framework", 1)
    add_section_intro(doc, "The audit determines whether each important page can enter the search and answer-engine pipeline, whether the content deserves selection, and whether visibility produces business value.")

    add_heading(doc, "Definitions", 2)
    add_table(
        doc,
        ["Term", "Working definition", "Primary evidence"],
        [
            ["SEO", "Improving how search engines discover, understand, index, present, and rank content so that people can find and act on it.", "Search Console, Bing Webmaster Tools, crawls, server logs, analytics, rankings, and conversions."],
            ["AEO", "Improving the probability that answer engines retrieve, cite, summarize, or recommend accurate information from the brand and its sources.", "AI feature reports, repeated prompt tests, citations, answer accuracy, referral analytics, and assisted outcomes."],
            ["GEO", "A related term focused on generative engine visibility. This framework treats GEO as part of the broader AEO workstream.", "The same pipeline measures used for AEO."],
            ["Entity", "A distinct real-world thing such as Bean Leaf Brew, a product, a location, a person, or a category, described consistently across owned and third-party sources.", "Site facts, structured data, Business Profile, Merchant Center, citations, directories, and reputable third-party pages."],
            ["Answer asset", "A page or page section designed to resolve a real user question with accurate, sourceable, and decision-ready information.", "Content review, query coverage, citations, engagement, and conversion paths."],
        ],
        widths=[0.85, 3.4, 2.6],
        font_size=8.1,
    )

    add_para(doc, "Google's current guidance treats AEO and GEO work for its generative search features as SEO, because those features rely on the Search index and quality systems. It also says there are no special schema types, mandatory content chunking rules, or Google benefits from llms.txt. This framework keeps a distinct AEO lens because crawler controls, AI citations, brand accuracy, and cross-platform reporting still require additional operational work. [1]")

    add_heading(doc, "The visibility pipeline", 2)
    add_numbered(doc, [
        "Access: The platform can request the page and its essential resources without a robots, firewall, authentication, or rate-limit block.",
        "Render: The platform can see the main content, links, product facts, images, and structured data in the rendered output.",
        "Index: The engine accepts the preferred canonical URL and understands its status, language, freshness, and relationship to other pages.",
        "Retrieve: The page is relevant enough to be selected for the user's query or a generated fan-out query.",
        "Select and cite: The system chooses the page as a source and exposes or attributes it in the answer.",
        "Absorb accurately: The answer uses the source's facts, evidence, and context without misrepresenting the brand.",
        "Engage and convert: The visibility creates a qualified visit, brand search, assisted purchase, direct purchase, inquiry, subscription, or repeat visit.",
    ])
    add_para(doc, "A page can fail at any stage. For example, a product page may be indexed but not retrieved because it lacks differentiated facts; it may be cited but drive no visits; or it may be used in an answer but with an outdated price. The audit records the first failing stage and the business impact.")

    add_heading(doc, "Evidence levels", 2)
    add_table(
        doc,
        ["Level", "Use in this framework", "Examples"],
        [
            ["High", "Direct platform requirements or first-party reporting. Implement unless a documented business constraint overrides it.", "Official crawler controls, Search Essentials, Merchant Center requirements, structured data policies, and first-party performance reports."],
            ["Medium", "Repeated observations supported by controlled or multi-platform studies. Test before broad rollout.", "Content structure that improves extractability, repeated prompt measurement, and comparative answer assets."],
            ["Exploratory", "A proposal or tactic with limited adoption or unclear causal effect. Use only as a low-cost experiment.", "llms.txt for services that consume it and unverified platform-specific formatting tactics."],
            ["Reject", "Manipulative, deceptive, unsupported, or high-risk work.", "Fake mentions, fabricated reviews, hidden prompt instructions, mass-generated query variants, schema that is not visible on the page, and guaranteed citation claims."],
        ],
        widths=[0.8, 3.0, 3.05],
        font_size=8.1,
    )

    add_heading(doc, "Audit scope and sampling", 2)
    add_bullets(doc, [
        "All indexable templates: home, shop, category or collection, product, about, contact, policy, editorial or brewing guide, campaign landing page, and any partner or wholesale pages.",
        "All traffic and revenue leaders, all strategic products, newly launched pages, pages with declining clicks, and pages already cited by AI systems.",
        "A representative sample of at least five URLs per template where volume allows, plus every URL with a critical status or manual action.",
        "Mobile and desktop rendering, public and logged-out access, multiple geographies where sales are targeted, and English plus any future language versions.",
        "Brand, category, product, comparative, informational, local, transactional, and post-purchase prompts across the major answer engines.",
        "At least three relevant competitors or reference publishers for search results, content depth, citations, entity signals, product data, and local prominence.",
    ])

    add_heading(doc, "Audit deliverables", 2)
    add_table(
        doc,
        ["Deliverable", "Minimum contents"],
        [
            ["Baseline score", "Weighted 100-point score, evidence links, confidence, and critical-failure overrides."],
            ["Issue register", "URL or template, issue, evidence, affected stage, business impact, severity, effort, owner, dependency, and recommended fix."],
            ["Content inventory", "Purpose, intent, topic or entity, lifecycle status, performance, quality, answer readiness, overlap, and next action."],
            ["Prompt benchmark", "Prompt, engine, location, run date, search activation, cited domains and URLs, brand mention, factual accuracy, sentiment, referral path, and screenshots or exports."],
            ["Competitor matrix", "Who appears, which source types win, recurring facts and formats, authority signals, product or local coverage, and defensible gaps."],
            ["Implementation backlog", "P0 to P3 priority, impact, effort, owner, acceptance criteria, validation method, release batch, and target date."],
        ],
        widths=[1.55, 5.3],
        font_size=8.3,
    )

    page_break(doc)
    add_heading(doc, "Audit Scorecard", 1)
    add_section_intro(doc, "The score provides structure, but a critical accessibility, indexation, security, policy, or data-integrity failure overrides the total.")

    score_rows = [
        ["Business outcomes and measurement", "5", "Goals, analytics integrity, conversions, revenue attribution, and baseline windows."],
        ["Crawler access and delivery", "10", "robots.txt, AI user agents, WAF or CDN access, status codes, logs, and essential resources."],
        ["Indexation and canonicalization", "12", "Index eligibility, canonical URLs, sitemaps, duplication, parameters, redirects, and freshness."],
        ["Rendering experience and accessibility", "10", "Mobile rendering, JavaScript output, semantic structure, Core Web Vitals, usability, and accessibility tree."],
        ["Architecture and internal links", "8", "Hierarchy, crawl depth, navigation, breadcrumbs, orphan pages, anchor text, and faceted paths."],
        ["On page search presentation", "8", "Titles, headings, snippets, images, open graph data, and result eligibility."],
        ["Content quality and intent coverage", "14", "Originality, usefulness, expertise, freshness, accuracy, completeness, and overlap."],
        ["Answer readiness and sourceability", "12", "Direct answers, comparisons, evidence, definitions, steps, FAQs, citations, and factual modularity."],
        ["Entities structured data and trust", "10", "Brand facts, authorship, policies, Organization, Product, Offer, LocalBusiness, and validation."],
        ["Authority and third party corroboration", "5", "Relevant links, reputable mentions, reviews, partnerships, directories, and community presence."],
        ["Local shopping and multimedia surfaces", "3", "Business Profile, Merchant Center, feeds, images, video, and local landing information."],
        ["AEO measurement and governance", "3", "Prompt panel, AI reports, brand accuracy review, testing discipline, ownership, and cadence."],
    ]
    add_table(doc, ["Dimension", "Weight", "What the score covers"], score_rows, widths=[2.3, 0.65, 3.9], font_size=8.0)

    add_heading(doc, "Scoring method", 2)
    add_table(
        doc,
        ["Rating", "Meaning", "Evidence standard"],
        [
            ["0", "Absent or seriously broken", "No control, repeated failure, or direct evidence of blocking or misleading output."],
            ["1", "Ad hoc", "Some work exists, but coverage is inconsistent and ownership is unclear."],
            ["2", "Partly implemented", "The main path works, but important templates, pages, or engines remain uncovered."],
            ["3", "Implemented", "Documented control with broad coverage, validation, and an owner."],
            ["4", "Measured and improving", "Complete control with trend data, alerting or scheduled review, and demonstrated improvement."],
        ],
        widths=[0.7, 2.35, 3.8],
        font_size=8.3,
    )
    add_para(doc, "For each dimension, calculate weighted points as rating divided by four multiplied by the dimension weight. Do not award points without evidence. Record confidence as high, medium, or low.")

    add_heading(doc, "Interpretation", 2)
    add_table(
        doc,
        ["Score", "Interpretation", "Action"],
        [
            ["85 to 100", "Strong foundation", "Protect quality, expand differentiated content, and run controlled AEO experiments."],
            ["70 to 84", "Usable but incomplete", "Fix coverage gaps and concentrate on high-value templates and measurable answer opportunities."],
            ["50 to 69", "Material constraints", "Prioritize technical, data, content, and entity foundations before scaling publication."],
            ["Below 50", "High risk", "Pause nonessential expansion and correct eligibility, integrity, and measurement failures."],
        ],
        widths=[1.0, 2.4, 3.45],
        font_size=8.4,
    )

    add_heading(doc, "Severity and priority", 2)
    add_table(
        doc,
        ["Priority", "Use when", "Examples"],
        [
            ["P0", "The issue blocks revenue pages, creates legal or security risk, exposes wrong information, or removes broad search eligibility.", "Sitewide noindex, broken checkout links, severe canonical error, product price mismatch, or essential bots blocked unintentionally."],
            ["P1", "The issue materially limits discovery, indexation, rich results, local or shopping eligibility, or high-value content performance.", "Missing product schema, orphan category pages, widespread duplicate titles, stale Business Profile, or critical Core Web Vitals failure."],
            ["P2", "The change improves relevance, source quality, answer coverage, usability, or measurement for an important segment.", "Upgrade product facts, publish brewing guides, improve comparison pages, add first-party evidence, or build prompt reporting."],
            ["P3", "The work is experimental, low impact, or cosmetic.", "Optional llms.txt trial, minor metadata polish, or low-volume long-tail page refinement."],
        ],
        widths=[0.75, 3.15, 2.95],
        font_size=8.0,
    )

    page_break(doc)
    add_heading(doc, "Technical SEO and Access Audit", 1)
    add_section_intro(doc, "This audit verifies that search and answer engines can request, render, index, and refresh the intended public content without exposing private or low-value URLs.")

    add_heading(doc, "Crawler and firewall access", 2)
    add_bullets(doc, [
        "Fetch robots.txt from every public hostname and subdomain. Confirm it returns 200, is not redirected unexpectedly, and does not block important pages, styles, scripts, images, product data, or canonical URLs.",
        "Compare robots rules with the intended policy for Googlebot, Bingbot, OAI-SearchBot, GPTBot, ChatGPT-User, PerplexityBot, Perplexity-User, ClaudeBot, Claude-SearchBot, and Claude-User. Search and training access must be decided separately where the platform supports it. [17] [19] [20] [21]",
        "Review CDN, web application firewall, bot management, authentication, geofencing, rate limits, and challenge pages. A robots allow rule does not help if the WAF returns 403, 429, a JavaScript challenge, or incomplete content.",
        "Use server logs to verify successful requests, response sizes, crawl frequency, last crawl, redirect paths, and error patterns. Verify declared bot IP ranges where a platform publishes them; do not trust the user agent alone.",
        "Check that private admin, account, checkout, order, internal search, preview, and API routes are protected appropriately. Do not rely on robots.txt to keep confidential content out of indexes.",
    ])

    add_heading(doc, "Status codes redirects and canonical URLs", 2)
    add_bullets(doc, [
        "Every indexable page should return a stable 200 response. Permanent moves should use one-hop 301 or 308 redirects. Avoid loops, chains, blanket redirects to the home page, and soft 404s.",
        "Deleted products should follow a documented rule: retain a useful replacement page when demand remains, redirect only to a genuinely equivalent product, or return 404 or 410 when no substitute exists.",
        "Each indexable page should declare a self-referencing canonical unless a justified consolidation rule applies. Canonical targets must be 200, indexable, equivalent, and internally linked.",
        "Normalize host, protocol, trailing slash, case, tracking parameters, sorting, filters, and duplicate variant paths. Confirm search engines select the intended canonical in inspection tools.",
        "Do not use canonical as a substitute for weak architecture or uncontrolled faceted navigation. Align redirects, canonicals, sitemap URLs, internal links, hreflang, and structured data URLs. [8]",
    ])

    add_heading(doc, "Sitemaps and freshness", 2)
    add_bullets(doc, [
        "Generate XML sitemaps containing only canonical, 200, indexable URLs. Segment by content type when it helps diagnosis, and keep each sitemap within protocol limits.",
        "Use accurate lastmod values only when meaningful page content changes. Remove deleted, redirected, noindex, or duplicate URLs promptly.",
        "Submit sitemaps in Google Search Console and Bing Webmaster Tools. Reconcile submitted, discovered, crawled, and indexed counts by template.",
        "For Bing and participating engines, evaluate IndexNow for product additions, price or availability changes, updates, and deletions. Submission improves discovery speed but does not guarantee indexation or ranking. [17] [18]",
    ])

    add_heading(doc, "Rendering and JavaScript", 2)
    add_bullets(doc, [
        "Compare raw HTML, rendered HTML, and the visible page. Critical titles, main copy, product facts, price, availability, links, image references, and structured data should be reliably available to crawlers.",
        "Test logged-out rendering with disabled cache, slow network, blocked third-party scripts, and JavaScript errors. Validate that hydration does not remove or replace key content.",
        "Ensure crawlable navigation uses real links with descriptive anchor text. Avoid click handlers or client-only controls as the only path to important pages.",
        "Verify lazy-loaded images and video have stable, discoverable URLs, descriptive alt text, relevant nearby text, and do not require user interaction before the primary asset exists in the rendered output.",
        "Check that product and offer data is server-delivered or otherwise reliable. Google specifically recommends Product structured data in initial HTML for merchant experiences when possible. [10]",
    ])

    add_heading(doc, "Mobile page experience and accessibility", 2)
    add_bullets(doc, [
        "Test mobile layouts, navigation, product variants, cart actions, form labels, error messages, tap targets, zoom, and keyboard operation.",
        "Measure field Core Web Vitals at the 75th percentile for mobile and desktop. Current good thresholds are LCP at or below 2.5 seconds, INP at or below 200 milliseconds, and CLS at or below 0.1. [14]",
        "Inspect the semantic outline and accessibility tree. Use one descriptive main heading, logical subheadings, named controls, text alternatives, table headers, and meaningful link text.",
        "Keep main content distinct from banners, overlays, cookie notices, ads, and repeated navigation. Avoid intrusive interstitials that delay access.",
        "Test visual and DOM interpretation for agentic browsing on critical product and contact tasks. Accessibility improvements also make browser-agent interaction more reliable. [1]",
    ])

    add_heading(doc, "Technical audit evidence", 2)
    add_table(
        doc,
        ["Test", "Evidence to retain", "Pass condition"],
        [
            ["Crawl and response", "Crawl export, log sample, status, final URL, response time, and response size.", "Important public URLs return consistent 200 responses and essential resources are accessible."],
            ["Index eligibility", "Robots, meta robots, X-Robots-Tag, canonical, sitemap, and inspection result.", "The intended canonical is indexable and eligible for a snippet."],
            ["Rendered content", "Raw versus rendered comparison and screenshots on mobile and desktop.", "Main content, links, media, and data are present and usable."],
            ["Performance", "Field data, lab traces, template distribution, and top causes.", "The majority of important URL groups meet current good thresholds or have a funded remediation plan."],
            ["AI access", "robots policy, WAF rules, verified logs, platform IP documentation, and owner approval.", "The implemented access policy matches the approved business decision."],
        ],
        widths=[1.4, 3.25, 2.2],
        font_size=8.0,
    )

    page_break(doc)
    add_heading(doc, "Architecture and On Page Audit", 1)
    add_section_intro(doc, "Architecture should show what the business sells, which pages matter, and how products, categories, guides, locations, and policies relate.")

    add_heading(doc, "Information architecture", 2)
    add_bullets(doc, [
        "Map the intended hierarchy from home to shop, category, product, guide, about, contact, shipping, returns, and support pages. Each indexable page should have one primary role.",
        "Keep important products and categories within a short, logical click path. Identify orphan pages and pages reachable only from sitemaps or internal search.",
        "Use consistent category definitions and avoid near-duplicate collections that compete for the same intent. For Bean Leaf Brew, separate product categories from editorial topics such as brewing methods, roast guidance, coffee origins, office premix, and gifting.",
        "Add breadcrumbs where they help users and engines understand hierarchy. Breadcrumb structured data must match the visible navigation path.",
        "Control sorting, filtering, pagination, and search-result URLs. Decide which combinations deserve indexable landing pages and which should remain crawlable but non-indexable, canonicalized, or blocked from crawling after careful testing.",
    ])

    add_heading(doc, "Internal linking", 2)
    add_bullets(doc, [
        "Link from high-authority and high-traffic pages to strategic categories, products, guides, and policies using descriptive anchor text.",
        "Connect products to relevant brewing guides, grind information, origin stories, serving guidance, comparison content, and compatible equipment. Link guides back to appropriate products without forcing a sales pitch.",
        "Use related-product and related-guide modules based on real relevance, not sitewide repetition. Check that links are server-rendered and accessible without interaction.",
        "Audit broken internal links, redirects in links, excessive repeated links, empty anchors, image links without alt text, and pages with unusually low or high internal-link counts.",
    ])

    add_heading(doc, "Titles headings and snippets", 2)
    add_bullets(doc, [
        "Give every indexable page a unique, concise title that accurately describes the page and distinguishes the product, category, use case, or location. Avoid boilerplate that consumes the useful part of the title.",
        "Align the HTML title, visible main heading, prominent page title, open graph title, and page topic. They need not be identical, but they should not conflict.",
        "Write a specific meta description for important pages. Treat it as a useful preview, not a ranking claim. Search engines may generate a different snippet from page content for a specific query. [3] [7]",
        "Use headings to explain the page in a logical order. Do not turn every sentence into a heading or create a repetitive question list solely for extraction.",
        "Use descriptive image filenames where practical, accurate alt text, relevant captions or nearby copy, multiple high-quality product angles, and stable image URLs.",
    ])

    add_heading(doc, "Template checks for Bean Leaf Brew", 2)
    add_table(
        doc,
        ["Template", "Primary intent", "Audit questions"],
        [
            ["Home", "Brand and category discovery", "Does it state what Bean Leaf Brew is, what it sells, who it serves, where it operates, why it is credible, and where users should go next?"],
            ["Shop and category", "Browse and compare", "Are categories distinct, indexable when valuable, internally linked, descriptive, filter-safe, and supported by useful selection guidance?"],
            ["Product", "Evaluate and purchase", "Are product name, images, price, availability, weight or variant, roast, notes, ingredients where relevant, preparation, shipping, returns, and proof complete and consistent?"],
            ["About", "Trust and entity understanding", "Does it establish the business identity, story, expertise, location, sourcing or roasting approach, team or founder where appropriate, and verifiable facts?"],
            ["Contact", "Local and support intent", "Are contact methods, service area, address if customer-facing, hours, response expectations, and map or profile links accurate?"],
            ["Policy", "Risk reduction and purchase confidence", "Are shipping, returns, refunds, privacy, terms, and cancellation rules public, current, linked from products and footer, and consistent with feeds and schema?"],
            ["Guide", "Learn and compare", "Does it solve one real task with expert input, original detail, clear steps or comparisons, cited sources where needed, and a natural path to products?"],
        ],
        widths=[1.1, 1.55, 4.2],
        font_size=7.8,
    )

    page_break(doc)
    add_heading(doc, "Content Quality and Intent Audit", 1)
    add_section_intro(doc, "The objective is a smaller, stronger body of content that answers real questions, demonstrates experience, and supports product and brand decisions.")

    add_heading(doc, "Content inventory fields", 2)
    add_table(
        doc,
        ["Field", "What to record"],
        [
            ["Identity", "URL, template, title, owner, publish date, last substantial review, language, and canonical."],
            ["Purpose", "Business goal, target audience, primary intent, funnel stage, target topic or entity, and desired action."],
            ["Performance", "Clicks, impressions, click rate, position, conversions, assisted conversions, revenue, backlinks, AI citations, and referral quality."],
            ["Quality", "Originality, first-hand experience, expert review, accuracy, evidence, completeness, readability, media, and freshness."],
            ["Overlap", "Competing URLs, duplicate sections, cannibalization, and opportunity to consolidate."],
            ["Action", "Keep, update, expand, consolidate, redirect, remove, noindex, or test."],
        ],
        widths=[1.15, 5.7],
        font_size=8.3,
    )

    add_heading(doc, "People first and source worthy review", 2)
    add_bullets(doc, [
        "Confirm the page has a clear intended audience and purpose beyond attracting search visits.",
        "Identify the original value: first-hand roasting or brewing experience, product testing, data, customer questions, process documentation, sourcing detail, comparisons, photographs, or expert explanation.",
        "Verify every material fact. Distinguish observed results from opinion, and update time-sensitive prices, stock, policies, locations, and claims.",
        "Show who created or reviewed the content when expertise matters. Link to a credible author, reviewer, or organization page and explain relevant experience without inflated credentials.",
        "Cite primary or authoritative sources for technical, health, safety, regulatory, or statistical claims. Do not add citations decoratively; each source should support the adjacent claim.",
        "Remove filler, copied summaries, repetitive introductions, and unsupported superlatives. There is no preferred word count; write enough to complete the task. [3] [4]",
        "Avoid scaled publication that produces many near-identical pages or fan-out query variants with little original value. Google treats scaled content created mainly to manipulate rankings or generative responses as spam. [5]",
    ])

    add_heading(doc, "Intent and topic coverage", 2)
    add_para(doc, "Build the content plan from customer decisions rather than a keyword list alone. For each topic, map the questions people ask before, during, and after purchase. Include both broad discovery and specific use cases.")
    add_table(
        doc,
        ["Intent", "Bean Leaf Brew examples", "Best page form"],
        [
            ["Brand", "What is Bean Leaf Brew, where is it based, what does it sell, how is it different?", "Home, about, contact, review and profile ecosystem."],
            ["Category", "Specialty coffee beans, ground coffee, instant premix, brewing equipment, coffee gifts.", "Category landing pages with selection guidance and product links."],
            ["Product", "Roast, origin, notes, grind, weight, price, availability, preparation, shelf life, delivery, and returns.", "Complete product detail pages with Product and Offer data."],
            ["Comparative", "Medium versus dark roast, beans versus ground coffee, premix versus brewed coffee, brew methods, or product comparisons.", "Neutral comparison guides with criteria, tradeoffs, and recommendations by need."],
            ["How to", "French press, pour over, moka pot, South Indian filter, storage, grind selection, or office preparation.", "First-hand brewing guides with steps, ratios, images or video, and troubleshooting."],
            ["Local", "Specialty coffee in Bengaluru, coffee delivery in Bengaluru, Whitefield coffee brand, office coffee supply.", "Accurate local and service-area content, Business Profile, contact and delivery information."],
            ["Transactional", "Buy coffee beans online, same-week delivery, gifting, subscriptions, wholesale, or office premix.", "Category, product, landing, policy, and contact pages with clear terms."],
            ["Post purchase", "How to brew this coffee, store it, adjust grind, manage a subscription, or resolve an order issue.", "Support pages, product-linked guides, email content, and structured help."],
        ],
        widths=[1.0, 3.4, 2.45],
        font_size=7.8,
    )

    add_heading(doc, "Content consolidation rules", 2)
    add_bullets(doc, [
        "Keep separate pages when they serve materially different intents, products, locations, or audiences and each has enough unique value.",
        "Merge pages when they compete for the same task and neither provides a complete answer. Preserve the strongest URL, redirect retired URLs, update links, and resubmit the sitemap.",
        "Update pages when the core purpose remains valid but facts, examples, media, or products have changed. Change the date only after a substantive review.",
        "Remove or noindex pages that have no useful audience, no unique value, and no strategic role. Check backlinks, revenue, and redirects before removal.",
    ])

    page_break(doc)
    add_heading(doc, "AEO and Answer Readiness Audit", 1)
    add_section_intro(doc, "Answer readiness is the ability of a page to supply an accurate, relevant, attributable part of a generated answer while remaining useful to a human reader.")

    add_heading(doc, "Answer asset criteria", 2)
    add_bullets(doc, [
        "The page resolves a real question or decision and states the answer early enough that a reader can confirm relevance.",
        "Key facts are explicit: product name, category, definitions, quantities, units, conditions, dates, service areas, prices, availability, and limitations.",
        "Comparisons use stable criteria and explain who each option suits. Tables are used when readers need to scan the same fields across options.",
        "Procedures include prerequisites, ordered steps, amounts or ratios, timing, expected result, and troubleshooting where relevant.",
        "Claims have evidence. The page links to first-party data, source material, methods, policies, or external authorities as appropriate.",
        "The page contains original information or perspective that cannot be recreated by summarizing the first search results.",
        "Content is internally consistent with product data, schema, feeds, Business Profile, support pages, images, video, and third-party listings.",
        "The source can be quoted or summarized without losing essential conditions. Avoid vague pronouns, unexplained labels, and claims whose subject is unclear.",
    ])

    add_heading(doc, "Question and passage review", 2)
    add_para(doc, "For each strategic page, list the five to fifteen questions it should be able to answer. Mark whether the page provides a complete answer, partial answer, conflicting answer, or no answer. Do not create a separate page or FAQ for every wording variation; improve the core page when the intent is the same.")
    add_table(
        doc,
        ["Review item", "Pass test", "Common failure"],
        [
            ["Definition", "The term is explained in plain language and tied to the relevant product or task.", "Circular definition or generic copy with no business context."],
            ["Fact", "The value has a unit, condition, date, and source or owner where needed.", "Price, weight, shelf life, or delivery statement differs across pages."],
            ["Comparison", "Options share clear criteria, tradeoffs, and recommendations by use case.", "One-sided promotional copy or mismatched criteria."],
            ["Procedure", "Steps are ordered, testable, and include quantities, timing, and expected result.", "Advice is inspirational but cannot be followed."],
            ["Recommendation", "The page explains the basis, audience, constraints, and commercial relationship.", "Unsupported best-product claim or hidden sponsorship."],
            ["Evidence", "Data, examples, reviews, or sources directly support the adjacent claim.", "Decorative citations or unverifiable numbers."],
        ],
        widths=[1.15, 3.0, 2.7],
        font_size=8.0,
    )

    add_heading(doc, "Platform presence testing", 2)
    add_para(doc, "Create a fixed prompt panel for each engine and location. Run prompts in clean, documented conditions. Record whether web search activated, which sources appeared, the order and prominence of citations, whether Bean Leaf Brew was mentioned, whether the answer was accurate, and whether a user could reach the intended page.")
    add_table(
        doc,
        ["Prompt group", "Examples to include", "What to measure"],
        [
            ["Brand accuracy", "What is Bean Leaf Brew, where is it based, what products does it sell, and how can someone contact it?", "Correct identity, description, location, links, sentiment, and unsupported claims."],
            ["Category discovery", "Which specialty coffee brands deliver in Bengaluru, or where can I buy medium roast coffee online in India?", "Mention rate, cited sources, competitors, product fit, and geographic accuracy."],
            ["Product and use case", "Coffee for French press, office premix, gifting, or South Indian filter preparation.", "Product inclusion, factual basis, price or stock accuracy, and landing-page match."],
            ["Comparative", "Beans versus ground coffee, roast choices, brewing equipment, or brand comparisons.", "Criteria, source diversity, citation share, recommendation rationale, and balance."],
            ["Informational", "Brew ratios, grind size, storage, tasting notes, or caffeine questions.", "Guide visibility, factual use, citation, accuracy, and follow-up coverage."],
            ["Local", "Specialty coffee near Whitefield or coffee delivery in Bengaluru.", "Business Profile visibility, distance sensitivity, reviews, NAP accuracy, and service-area clarity."],
        ],
        widths=[1.25, 3.5, 2.1],
        font_size=7.8,
    )

    add_heading(doc, "Measurement reliability", 2)
    add_bullets(doc, [
        "Repeat prompts across runs, paraphrases, dates, engines, devices, and locations. A single answer is an observation, not a benchmark.",
        "Retain non-search and non-citation outputs in the denominator. Report search activation, mention rate, and citation rate separately.",
        "Canonicalize URLs when calculating share so tracking parameters, anchors, redirects, and translated variants do not fragment the source.",
        "Use human review for brand accuracy, claim support, and recommendation quality. Automated judges can assist but should not be the only measure.",
        "The 2026 research literature finds substantial variation across engines and time and does not establish a stable cross-platform tactic that guarantees discoverability, traffic, or revenue. Treat citations as an intermediate signal and test downstream effects. [24]",
    ])

    add_heading(doc, "AEO audit outputs", 2)
    add_table(
        doc,
        ["Output", "Definition"],
        [
            ["Prompt coverage", "Percentage of the approved prompt universe tested in the period."],
            ["Search activation rate", "Percentage of runs in which the platform used web retrieval or exposed sources."],
            ["Brand mention rate", "Percentage of all runs that mention Bean Leaf Brew, whether cited or not."],
            ["Citation rate", "Percentage of all runs that cite at least one canonical Bean Leaf Brew URL."],
            ["Citation share", "Bean Leaf Brew citations divided by all citations in the tested answer set, with engine and prompt group shown separately."],
            ["Accuracy rate", "Percentage of reviewed brand claims that are correct, current, and supported."],
            ["Referral quality", "Engagement, conversion, revenue, assisted conversion, and new-customer quality for identifiable AI referrals."],
        ],
        widths=[1.7, 5.15],
        font_size=8.2,
    )

    page_break(doc)
    add_heading(doc, "Entity Structured Data and Trust Audit", 1)
    add_section_intro(doc, "Entity and structured data work should clarify real information that is already visible and consistent, not manufacture authority.")

    add_heading(doc, "Brand entity consistency", 2)
    add_bullets(doc, [
        "Create an approved fact sheet for legal name, trading name, logo, website, primary category, description, founding information, contact details, customer-facing address if applicable, service area, social profiles, support channels, and policy URLs.",
        "Compare the fact sheet with the home page, about page, contact page, footer, Organization markup, Google Business Profile, Merchant Center, social profiles, directories, partner pages, packaging, invoices, and media coverage.",
        "Resolve material contradictions. A single consistent phrase is not required, but names, locations, ownership, products, and policies must agree.",
        "Maintain clear authorship or reviewer information for expert content. Do not invent credentials or use generic team bylines when a specific person is responsible.",
        "Audit third-party knowledge sources and high-ranking pages for incorrect brand facts. Prioritize corrections by visibility and risk.",
    ])

    add_heading(doc, "Structured data rules", 2)
    add_bullets(doc, [
        "Use JSON-LD or another supported format. Select the most specific applicable types and include required and useful recommended properties.",
        "Markup must describe the page it appears on and match visible content. Do not mark hidden, misleading, unavailable, or unrelated information.",
        "Use stable identifiers and URLs, including @id references where they help connect Organization, WebSite, Product, Offer, policy, and page entities.",
        "Validate syntax with Schema.org Validator and eligibility with Google's Rich Results Test where applicable. Then test rendered URLs and monitor Search Console enhancement reports.",
        "Correct warnings that reflect meaningful missing information. Do not add fields only to silence a tool when the information is not available or visible.",
        "Structured data can improve understanding and rich-result eligibility, but Google does not guarantee a rich result, and no special schema is required for its generative AI features. [1] [9]",
    ])

    add_heading(doc, "Recommended schema coverage", 2)
    add_table(
        doc,
        ["Page", "Recommended types", "Key checks"],
        [
            ["Home or about", "Organization or the most specific appropriate subtype, WebSite", "Name, URL, logo, contact, address or service area where applicable, sameAs, policies, and stable @id."],
            ["Contact or location", "Organization and LocalBusiness when a real customer-facing business location qualifies", "Address, telephone, email, opening hours, map, area served, and consistency with Business Profile."],
            ["Product", "Product, Offer, Brand, BreadcrumbList", "Name, images, description, SKU or identifier where used, brand, price, INR currency, availability, URL, condition, shipping and return policy references, and variant logic."],
            ["Category", "CollectionPage, ItemList when useful, BreadcrumbList", "Do not expect product rich results from a general listing page; ensure listed items and URLs match the visible page."],
            ["Guide or article", "Article or BlogPosting, Person or Organization author, BreadcrumbList", "Headline, author, date published, date modified after substantive review, image, mainEntityOfPage, and citations in visible copy."],
            ["Recipe when applicable", "Recipe, VideoObject or ImageObject where supported", "Use only for a genuine recipe with visible ingredients, quantities, steps, time, yield, and media."],
            ["FAQ section", "FAQPage only when the page genuinely contains visible questions and answers", "Do not add repetitive questions for markup. Do not assume rich-result eligibility; Google limits FAQ rich results."],
        ],
        widths=[1.25, 2.15, 3.45],
        font_size=7.55,
    )

    add_heading(doc, "Product and merchant data", 2)
    add_bullets(doc, [
        "Ensure the page, Product and Offer markup, Merchant Center feed, checkout, and inventory system agree on product identity, variant, price, currency, availability, image, landing URL, shipping, and return policy.",
        "Use Merchant Center alongside on-page structured data. Google states that using both maximizes eligibility and helps it understand and verify product data. [10]",
        "Provide high-quality crawlable product images, including useful aspect ratios and product-only views. Avoid text-heavy promotional overlays in primary feed images.",
        "Define variant identity and canonical rules for weight, grind, equipment, premix flavor, pack size, and other options. Do not create indexable URLs for every parameter unless each version is stable, discoverable, and useful.",
        "Publish complete shipping and return information. India is among the countries where Merchant Center requires shipping settings or shipping costs in product data for free listings. [27]",
    ])

    add_heading(doc, "Local trust and prominence", 2)
    add_bullets(doc, [
        "Verify and complete the Google Business Profile if the business qualifies. Keep category, description, address or service area, hours, phone, website, photos, products, and special hours accurate.",
        "Maintain consistent name, address, and phone information across owned pages and reputable directories. Remove duplicates and correct outdated profiles.",
        "Develop a lawful review process that asks real customers for honest feedback without gating, incentives that violate platform rules, or fabricated text. Respond to reviews with useful, privacy-safe information.",
        "Create local pages only when they represent a real location, service area, or distinct customer need. Avoid doorway pages that swap city names without unique value.",
        "Google describes local visibility mainly through relevance, distance, and prominence. Distance cannot be optimized away; improve complete information, relevance, reviews, and real-world prominence. [26]",
    ])

    page_break(doc)
    add_heading(doc, "Authority and Reputation Audit", 1)
    add_section_intro(doc, "The audit looks for genuine corroboration and reputation signals, not a volume of low-quality links or mentions.")

    add_heading(doc, "Link and mention quality", 2)
    add_bullets(doc, [
        "Inventory linking domains, unlinked brand mentions, reviews, partner pages, supplier or marketplace pages, event pages, local directories, interviews, media coverage, and community references.",
        "Prioritize topical relevance, editorial independence, audience value, source reputation, and accurate context. One strong relevant source can matter more than many generic placements.",
        "Review anchor text and destination quality. Branded and natural descriptive links should point to the most useful canonical page.",
        "Identify toxic patterns such as paid link networks, irrelevant directories, automated guest posts, hidden links, sitewide keyword anchors, and fabricated reviews. Do not pursue removal work solely because a third-party tool labels a link risky; investigate context.",
        "Track inaccurate high-visibility descriptions, outdated prices, wrong locations, or obsolete product details and request corrections from the source where appropriate.",
    ])

    add_heading(doc, "Earned authority opportunities", 2)
    add_table(
        doc,
        ["Opportunity", "Useful asset", "Integrity rule"],
        [
            ["Coffee education", "Original brewing tests, grind charts, water and ratio experiments, tasting methods, or roast explanations.", "Publish methods and limits; do not convert opinions into scientific claims."],
            ["Local relevance", "Bengaluru coffee guides, community events, verified stockists, local collaborations, and delivery information.", "Use real participation and accurate locations, not city-page duplication."],
            ["Product proof", "Transparent specifications, sourcing details, images, batch or roast information, preparation guidance, and verified reviews.", "Keep claims current and consistent with packaging and feeds."],
            ["Business use cases", "Office premix setup, hospitality supply, gifting, or wholesale documentation with clear service terms.", "Disclose commercial relationships and avoid unsupported savings or performance claims."],
            ["Partnerships", "Supplier, cafe, event, creator, or nonprofit collaboration pages.", "Earn editorial references through real work; do not buy inauthentic mentions."],
        ],
        widths=[1.2, 3.5, 2.15],
        font_size=7.9,
    )

    add_heading(doc, "Reputation review", 2)
    add_bullets(doc, [
        "Review search and answer results for the brand, founders or spokespeople where public, key products, reviews, complaints, returns, delivery, and customer service.",
        "Classify each issue as factual error, outdated information, legitimate criticism, service failure, impersonation, policy violation, or low-value noise.",
        "Correct owned information first. Respond to legitimate public feedback with facts and resolution paths. Escalate impersonation, fraud, or unsafe content through platform and legal processes.",
        "Measure improvement through accuracy and issue resolution, not by attempting to suppress legitimate criticism.",
    ])

    page_break(doc)
    add_heading(doc, "Measurement and Governance Audit", 1)
    add_section_intro(doc, "Measurement must connect platform visibility to qualified demand and business outcomes while preserving a repeatable evidence trail.")

    add_heading(doc, "Measurement stack", 2)
    add_table(
        doc,
        ["Source", "Use", "Audit checks"],
        [
            ["Google Search Console", "Queries, pages, countries, devices, indexation, rich results, Core Web Vitals, and generative AI visibility reports.", "Verified properties, correct users, sitemap status, annotations, exports, and consistent reporting windows."],
            ["Bing Webmaster Tools", "Search, crawling, IndexNow, site diagnostics, and AI Performance citations and grounding topics.", "Verified site, sitemap and IndexNow coverage, citation pages, query topics, and trend exports."],
            ["Analytics", "Organic and identifiable AI referrals, landing pages, engagement, conversions, assisted conversions, revenue, and cohorts.", "Channel definitions, referral exclusions, UTM policy, consent, ecommerce events, duplicate events, and cross-domain behavior."],
            ["Server and CDN logs", "Crawler access, errors, frequency, response size, WAF blocks, and freshness.", "Retention, bot verification, privacy controls, parsing rules, and alert thresholds."],
            ["Rank and prompt panel", "Traditional visibility and repeated AI answer observations across engines and locations.", "Stable query set, run protocol, canonical URLs, screenshots or exports, and human accuracy review."],
            ["Business systems", "Orders, revenue, margin, new customers, subscriptions, wholesale leads, support contacts, and repeat purchase.", "Join keys, time windows, attribution limits, and privacy-safe reporting."],
        ],
        widths=[1.4, 3.05, 2.4],
        font_size=7.6,
    )

    add_para(doc, "As of 2026, Google provides dedicated generative AI performance reporting for visibility in AI Overviews, AI Mode, and related experiences, and Bing provides AI Performance reporting for citations and grounding queries. These first-party reports should anchor the dashboard where available. [15] [18]")

    add_heading(doc, "Attribution rules", 2)
    add_bullets(doc, [
        "Report search visibility, AI visibility, referral traffic, and revenue as separate layers. Do not present citation growth as revenue growth.",
        "Use first-click, last-click, and assisted views for major decisions. AI exposure may create later brand searches or direct visits that are not visible as an AI referral.",
        "Annotate product launches, site releases, promotions, outages, feed changes, algorithm updates, and major media activity.",
        "Compare like periods and segment by template, query or prompt group, device, country, new versus returning customer, and branded versus non-branded demand.",
        "State data limitations in every executive view, including missing query data, platform sampling, attribution windows, personalization, and answer variability.",
    ])

    add_heading(doc, "Governance checks", 2)
    add_bullets(doc, [
        "A named executive or business owner approves objectives, risk posture, and budget.",
        "SEO, content, engineering, ecommerce, brand, analytics, customer service, legal, and security responsibilities are documented.",
        "Changes to robots, canonicals, redirects, schema, feeds, page templates, and analytics have code review, testing, rollback, and monitoring requirements.",
        "Content has a documented brief, factual owner, reviewer, publication criteria, substantial-update rule, and refresh or retirement date.",
        "AI crawler access and model-training preferences have approved owners and are reviewed at least quarterly or after material platform changes.",
        "No vendor may promise guaranteed rankings or AI citations, claim access to private platform metrics, or publish automated content without review. [25]",
    ])

    page_break(doc)
    add_heading(doc, "Part Two Implementation Framework", 1)
    add_section_intro(doc, "Implementation follows dependencies: establish measurement and policy, remove technical blockers, strengthen high-value pages and product data, then expand answer assets and authority.")

    add_heading(doc, "Implementation principles", 2)
    add_numbered(doc, [
        "Protect revenue and trust first. Fix wrong prices, stock, policies, broken purchase paths, privacy risks, and accidental indexation or blocking before growth work.",
        "Work by template and intent, not isolated pages. Correct shared causes in the application, data model, feed, or editorial process, then validate representative URLs.",
        "Release in controlled batches. Test a small set, inspect search and analytics evidence, and expand only when acceptance criteria are met.",
        "Keep visible content, structured data, feeds, and third-party profiles synchronized from authoritative business data.",
        "Build original sources before optimizing presentation. Formatting cannot compensate for thin, copied, or unsupported content.",
        "Measure leading and lagging indicators. Eligibility and citations may change before traffic, demand, or revenue.",
        "Maintain reversibility. Every technical or content release needs an owner, change log, validation plan, and rollback path.",
    ])

    add_heading(doc, "Definition of done", 2)
    add_table(
        doc,
        ["Work type", "Done means"],
        [
            ["Technical", "The fix is deployed across the intended scope, crawled in test and production, verified in rendered output, covered by automated checks where feasible, and monitored."],
            ["Content", "The page satisfies the approved brief, contains verified facts and original value, has an owner and review date, links correctly, and passes editorial and accessibility review."],
            ["Structured data", "Markup matches visible content, validates, uses stable IDs and canonical URLs, is present in rendered output, and is monitored in platform reports."],
            ["Feed or profile", "The authoritative data source is correct, the platform accepts it, diagnostics are resolved, and the public presentation matches the site."],
            ["AEO experiment", "Hypothesis, prompt set, engines, variants, dates, repetitions, success metric, guardrails, results, and decision are recorded."],
        ],
        widths=[1.4, 5.45],
        font_size=8.2,
    )

    add_heading(doc, "Impact and effort prioritization", 2)
    add_para(doc, "Prioritize by expected business impact, number and value of affected URLs, severity, confidence in the diagnosis, implementation effort, dependencies, and time to learn. Do not prioritize a task only because a tool labels it an error.")
    add_table(
        doc,
        ["Priority pattern", "Typical work"],
        [
            ["High impact low effort", "Correct metadata, remove accidental noindex, fix sitemap errors, align product facts, add missing policies, repair internal links, or update profile data."],
            ["High impact high effort", "Rebuild category architecture, improve rendering, implement product feeds and schema, redesign product data, or create a durable guide library."],
            ["Low impact low effort", "Test minor snippet improvements or optional emerging files after foundations are stable."],
            ["Low impact high effort", "Defer speculative platform hacks, bulk low-demand pages, or complex markup without visible user value."],
        ],
        widths=[1.65, 5.2],
        font_size=8.3,
    )

    page_break(doc)
    add_heading(doc, "Ninety Day Roadmap", 1)
    add_section_intro(doc, "The exact sequence should follow the audit, but this default roadmap fits an ecommerce and local brand with a modern rendered site.")

    add_heading(doc, "Days 0 to 15 Baseline and controls", 2)
    add_bullets(doc, [
        "Confirm business goals, priority product categories, target geographies, conversion events, revenue data, and baseline periods.",
        "Verify Google Search Console, Bing Webmaster Tools, analytics, Merchant Center, Business Profile, and log access. Capture current exports.",
        "Approve the crawler and AI-training policy. Check robots.txt, WAF, CDN, and server logs for Google, Bing, OpenAI, Perplexity, and Anthropic agents.",
        "Crawl the public site and inventory indexable URLs, templates, status codes, canonicals, metadata, headings, internal links, images, and structured data.",
        "Build the content inventory, page-value map, competitor set, and prompt universe. Run the first multi-engine benchmark.",
        "Create the issue register, weighted score, P0 and P1 list, and release governance.",
    ])

    add_heading(doc, "Days 16 to 30 Eligibility and data integrity", 2)
    add_bullets(doc, [
        "Resolve P0 crawl, index, redirect, canonical, rendering, security, or data conflicts.",
        "Publish or correct XML sitemaps and accurate lastmod values. Implement IndexNow where approved and useful.",
        "Fix global titles, descriptions, main headings, open graph data, image access, and internal-link defects on priority templates.",
        "Publish complete shipping, return, refund, cancellation, privacy, and contact information and link it from products and footer.",
        "Create the approved organization fact sheet and align the site, Business Profile, Merchant Center, and major third-party profiles.",
        "Implement baseline Organization, WebSite, Product, Offer, and Breadcrumb structured data with validation and monitoring.",
    ])

    add_heading(doc, "Days 31 to 60 High value page upgrades", 2)
    add_bullets(doc, [
        "Rewrite or enrich priority product pages with differentiated descriptions, specifications, preparation, variants, availability, delivery, returns, original imagery, and relevant FAQs.",
        "Create or improve category pages with selection guidance, distinct intent, crawlable product links, and safe filter handling.",
        "Upgrade home, about, and contact pages so the business identity, location, expertise, products, policies, and proof are explicit.",
        "Launch the first answer assets based on evidence: brewing guides, roast or grind selection, storage, product comparisons, office premix guidance, and local delivery information.",
        "Connect products, categories, and guides with contextual internal links and breadcrumbs.",
        "Resolve Merchant Center diagnostics and confirm feed and page parity for price, stock, image, shipping, return, variant, and URL data.",
    ])

    add_heading(doc, "Days 61 to 90 Authority and optimization", 2)
    add_bullets(doc, [
        "Publish two to four evidence-led content assets with original tests, expert input, images or video, and clear methods.",
        "Run legitimate digital PR, partner, community, supplier, stockist, or creator outreach around the original assets and real business activity.",
        "Improve Business Profile media, products, service details, review response process, and local consistency.",
        "Repeat the prompt benchmark and compare brand accuracy, search activation, mention rate, citations, source mix, and referral quality.",
        "Test selected title, introduction, answer passage, comparison, media, and internal-link changes. Keep tests scoped and preserve control pages or time windows.",
        "Review scorecard progress, publish the next quarterly roadmap, and assign refresh dates for priority content and data feeds.",
    ])

    add_heading(doc, "Expected ninety day outputs", 2)
    add_table(
        doc,
        ["Area", "Expected output"],
        [
            ["Foundation", "Approved crawl policy, clean sitemap set, stable canonicals, indexation map, working monitoring, and resolved critical errors."],
            ["Commerce", "Accurate product facts, Product and Offer markup, Merchant Center parity, strong images, and complete shipping and return information."],
            ["Content", "Upgraded priority product and category pages plus a first set of source-worthy brewing, comparison, local, or business-use content."],
            ["Entity and local", "Consistent brand fact sheet, Organization and LocalBusiness coverage where eligible, accurate Business Profile, and correction queue for third-party facts."],
            ["AEO", "Approved prompt panel, baseline and follow-up runs, brand accuracy log, citation map, and controlled experiment backlog."],
            ["Operations", "Named owners, acceptance criteria, release checklist, dashboard, monthly review, quarterly strategy cycle, and content refresh calendar."],
        ],
        widths=[1.35, 5.5],
        font_size=8.2,
    )

    page_break(doc)
    add_heading(doc, "Implementation Workstreams", 1)
    add_section_intro(doc, "Each workstream has a distinct output, but the work should be planned together so pages, data, feeds, and measurement stay synchronized.")

    add_heading(doc, "Technical search foundation", 2)
    add_numbered(doc, [
        "Define indexable URL rules for every route and template, including product variants, filters, internal search, checkout, admin, and tracking parameters.",
        "Implement consistent 200, redirect, 404, and 410 behavior with automated tests for critical paths.",
        "Align canonical tags, sitemap URLs, internal links, structured data URLs, and feed landing URLs.",
        "Ensure critical content and metadata render reliably for crawlers and logged-out users. Monitor JavaScript and hydration failures.",
        "Create bot-specific access tests and alert on unexpected 403, 429, challenge, or material response-size changes.",
        "Track Core Web Vitals and template performance. Fix image weight, server response, layout shifts, long tasks, third-party scripts, and cache behavior in impact order.",
        "Add deployment checks for robots, noindex, canonical, sitemap, schema, status codes, and analytics events.",
    ])

    add_heading(doc, "Content and information architecture", 2)
    add_numbered(doc, [
        "Define the topic and product architecture, including one primary page for each important user intent.",
        "Create briefs from customer questions, search evidence, support logs, sales conversations, product data, and competitor gaps.",
        "Require original value in every new page: a test, process, expert explanation, data point, comparison, image set, video, or business-specific answer.",
        "Write descriptive titles and headings, answer the main task directly, and use comparison tables or steps only when they improve decisions.",
        "Add authorship, review, sources, update dates, and limitations where the topic requires them.",
        "Connect the page to relevant products, categories, policies, and related guides with descriptive internal links.",
        "Assign a review date based on volatility: product data and policy pages more frequently, evergreen education less frequently unless facts change.",
    ])

    add_heading(doc, "Answer assets", 2)
    add_para(doc, "Build answer assets only where Bean Leaf Brew can contribute useful, defensible information. A recommended initial portfolio follows.")
    add_table(
        doc,
        ["Asset", "Required substance", "Primary outcome"],
        [
            ["Coffee selection guide", "Roast, flavor, brew method, grind, milk use, intensity, budget, and product fit with clear tradeoffs.", "Category discovery and decision support."],
            ["Brewing guides", "Equipment, dose, water, temperature, grind, steps, time, expected result, adjustment rules, images or video, and linked products.", "Informational visibility and post-purchase success."],
            ["Product comparison", "Stable criteria, audience, differences, limitations, pricing context, and recommendation by need.", "Comparative search and generated answers."],
            ["Office premix guide", "Serving assumptions, preparation, storage, equipment, ordering, service area, and contact path.", "Business and wholesale leads."],
            ["Bengaluru delivery guide", "Service area, timelines, cutoff, charges, product availability, support, and current limitations.", "Local and transactional visibility."],
            ["Coffee glossary", "Concise definitions linked to deeper guides and products where useful.", "Entity clarity and beginner discovery."],
            ["Original tests", "Method, variables, observations, photographs, results, limits, and reproducible guidance.", "Source worthiness, citations, and earned links."],
        ],
        widths=[1.45, 3.75, 1.65],
        font_size=7.8,
    )

    add_heading(doc, "Structured data merchant and local", 2)
    add_numbered(doc, [
        "Create one authoritative product-data model and map it to visible pages, Product and Offer JSON-LD, Merchant Center, checkout, email, and inventory.",
        "Implement Organization and WebSite entities with stable identifiers. Add LocalBusiness only when the business and location meet the platform's requirements.",
        "Implement Product and Offer data for each purchasable product and supported variant pattern, including INR price, availability, images, brand, shipping, and return references.",
        "Validate in development, production, Schema.org Validator, Rich Results Test, URL Inspection, Search Console, and Merchant Center diagnostics.",
        "Define Business Profile ownership, update procedures, photo cadence, review response standards, and holiday-hour controls.",
        "Reconcile public facts monthly and after any product, policy, location, contact, or logistics change.",
    ])

    add_heading(doc, "Authority and distribution", 2)
    add_numbered(doc, [
        "Turn original research, product launches, partnerships, workshops, events, and community activity into accurate source pages before outreach.",
        "Create a target list of relevant local, coffee, food, hospitality, office, ecommerce, and business publications or communities.",
        "Offer useful data, demonstrations, expert commentary, samples for transparent review, or collaboration. Disclose commercial relationships.",
        "Maintain supplier, partner, stockist, marketplace, directory, and social profiles with correct brand and product information.",
        "Track earned links, citations, unlinked mentions, profile accuracy, referral quality, and corrections. Avoid link-volume targets.",
    ])

    add_heading(doc, "Measurement and experimentation", 2)
    add_numbered(doc, [
        "Create a dashboard with eligibility, traditional search, generative search, citation, accuracy, referral, conversion, and revenue layers.",
        "Maintain a versioned prompt panel with stable core prompts and a smaller rotating discovery set.",
        "Repeat runs and paraphrases on a fixed schedule, and separate results by engine, location, prompt group, and search activation.",
        "Use one primary hypothesis per experiment and define guardrails for factual accuracy, organic traffic, conversion, and user experience.",
        "Prefer controlled page groups, staggered releases, or matched time windows. Record platform changes and external campaigns.",
        "Stop or reverse changes that harm factual accuracy, conversion, indexation, page experience, or policy compliance, even if citation counts rise.",
    ])

    page_break(doc)
    add_heading(doc, "Operating Model", 1)
    add_section_intro(doc, "SEO and AEO become durable only when routine business changes update the site, feeds, profiles, and measurement together.")

    add_heading(doc, "Ownership matrix", 2)
    add_table(
        doc,
        ["Area", "Accountable owner", "Required contributors"],
        [
            ["Objectives and budget", "Business owner or growth lead", "Ecommerce, finance, brand, and analytics."],
            ["Technical search", "Engineering lead", "SEO, security, infrastructure, QA, and analytics."],
            ["Product data and feeds", "Ecommerce or catalog owner", "Engineering, operations, finance, logistics, and SEO."],
            ["Content quality", "Content or brand lead", "Subject expert, SEO, design, legal where needed, and customer service."],
            ["Local profiles and reviews", "Operations or customer experience lead", "Brand, local staff, support, and SEO."],
            ["AI crawler policy", "Business owner with legal and security", "Engineering, SEO, content, and data governance."],
            ["Measurement", "Analytics owner", "SEO, ecommerce, finance, engineering, and business owner."],
            ["Reputation and corrections", "Brand or communications lead", "Customer service, legal, SEO, and operations."],
        ],
        widths=[1.55, 2.2, 3.1],
        font_size=7.9,
    )

    add_heading(doc, "Operating cadence", 2)
    add_table(
        doc,
        ["Frequency", "Review"],
        [
            ["Continuous", "Production errors, uptime, robots changes, WAF blocks, schema failures, feed disapprovals, checkout issues, and analytics outages."],
            ["Weekly", "Priority issue register, releases, product-data conflicts, content production, Merchant Center diagnostics, and major search anomalies."],
            ["Monthly", "Search and AI visibility, landing-page performance, prompt panel, citations, brand accuracy, Business Profile, reviews, and referral conversions."],
            ["Quarterly", "Full scorecard, crawler policy, competitor and source review, topic portfolio, authority plan, content refresh and retirement, and next roadmap."],
            ["After material change", "Domain, route, rendering, navigation, product data, policy, location, tracking, CDN, or platform change receives targeted re-audit."],
        ],
        widths=[1.2, 5.65],
        font_size=8.2,
    )

    add_heading(doc, "Release checklist", 2)
    add_bullets(doc, [
        "Acceptance criteria and affected URLs or templates are documented.",
        "Preview and production URLs are tested for status, robots, noindex, canonical, rendered content, internal links, metadata, schema, analytics, accessibility, and mobile layout.",
        "Sitemap, feed, profile, and policy dependencies are updated where relevant.",
        "A monitoring window, owner, rollback method, and expected lag are recorded.",
        "Search Console, Bing, Merchant Center, analytics, logs, and prompt benchmarks are checked after the release according to expected processing time.",
    ])

    add_heading(doc, "Editorial checklist", 2)
    add_bullets(doc, [
        "The page has a specific audience, task, and desired action.",
        "The main answer is clear and the page adds original value.",
        "Product, policy, location, and business facts are verified by their owners.",
        "Claims have appropriate evidence and commercial relationships are disclosed.",
        "Title, main heading, URL, description, images, alt text, links, author, dates, and structured data are complete.",
        "The page does not duplicate or cannibalize a stronger page.",
        "Mobile, accessibility, rendering, and conversion paths are reviewed.",
        "A refresh or retirement trigger is assigned.",
    ])

    page_break(doc)
    add_heading(doc, "Part Three Practical Reference", 1)
    add_section_intro(doc, "These reference patterns support implementation and should be adapted to the approved strategy, platform rules, and actual business data.")

    add_heading(doc, "Page content patterns", 2)
    add_table(
        doc,
        ["Page", "Recommended content order"],
        [
            ["Product", "Descriptive name and summary; image gallery; price, currency, stock and variants; purchase control; key specifications; taste or use details; preparation; delivery and returns; proof and reviews; related guides; related products; product-specific questions."],
            ["Category", "Clear category definition; who it suits; selection criteria; crawlable product list; concise comparison guidance; relevant guide links; delivery or policy note; category questions only when genuinely useful."],
            ["Comparison", "Decision and audience; comparison criteria; side-by-side table; evidence and method; differences and tradeoffs; recommendation by use case; limitations; product and guide links; update date."],
            ["Brewing guide", "Outcome; equipment and ingredients; ratio and variables; numbered steps; expected result; adjustment rules; common problems; images or video; source or tester; relevant products; review date."],
            ["Local landing", "Real local purpose; service or location facts; delivery area and timing; hours and contact; products or services; directions when applicable; local proof and imagery; policies; Business Profile link."],
            ["About", "Plain business description; origin and purpose; expertise and process; products and customers; location and service area; people where appropriate; verifiable proof; policies and contact; external profiles."],
        ],
        widths=[1.3, 5.55],
        font_size=8.0,
    )

    add_heading(doc, "Crawler policy decision matrix", 2)
    add_table(
        doc,
        ["Agent", "Primary purpose", "Default planning position", "Audit note"],
        [
            ["Googlebot", "Google Search and its generative search features", "Allow public indexable content", "Google states its AI search features use the Search index and Googlebot controls. Snippet controls affect direct use in AI features. [1] [7]"],
            ["Bingbot", "Bing Search, Copilot grounding, and partner experiences", "Allow public indexable content", "Check Bing Webmaster Tools, IndexNow, and WAF access. [17]"],
            ["OAI-SearchBot", "ChatGPT search results", "Allow if ChatGPT search visibility is desired", "OpenAI separates search access from GPTBot training access. [19]"],
            ["GPTBot", "Potential foundation-model training", "Business, legal, and content decision", "Allowing search does not require allowing training. [19]"],
            ["ChatGPT-User", "User-initiated retrieval", "Business, security, and content decision", "OpenAI says robots rules may not apply to user-initiated actions; verify other controls. [19]"],
            ["PerplexityBot", "Perplexity search indexing and links", "Allow if Perplexity visibility is desired", "Allow published IP ranges through the WAF when approved. [20]"],
            ["Perplexity-User", "User-initiated retrieval", "Business, security, and content decision", "Perplexity states this fetcher generally ignores robots.txt because the fetch is user requested. [20]"],
            ["Claude-SearchBot", "Claude search quality and indexing", "Allow if Claude search visibility is desired", "Anthropic documents separate search, user, and training robots. [21]"],
            ["ClaudeBot", "Potential model training", "Business, legal, and content decision", "Review training preference separately from search access. [21]"],
            ["Claude-User", "User-initiated retrieval", "Business, security, and content decision", "Disabling can reduce visibility in user-directed web search. [21]"],
        ],
        widths=[1.2, 1.65, 1.85, 2.15],
        font_size=7.15,
    )

    add_heading(doc, "Illustrative robots policy", 2)
    add_para(doc, "The following is a planning example, not a ready-to-deploy file. Paths, hostnames, legal decisions, and WAF controls must be validated before release.")
    example = (
        "User-agent: *\n"
        "Disallow: /admin/\n"
        "Disallow: /api/\n"
        "Disallow: /checkout/\n"
        "Disallow: /account/\n\n"
        "User-agent: OAI-SearchBot\n"
        "Allow: /\n\n"
        "User-agent: GPTBot\n"
        "Disallow: /\n\n"
        "User-agent: PerplexityBot\n"
        "Allow: /\n\n"
        "User-agent: Claude-SearchBot\n"
        "Allow: /\n\n"
        "User-agent: ClaudeBot\n"
        "Disallow: /\n\n"
        "Sitemap: https://www.example.com/sitemap.xml"
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.right_indent = Inches(0.35)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    for idx, line in enumerate(example.splitlines()):
        r = p.add_run(line)
        set_run_font(r, name="Consolas", size=8.2)
        if idx < len(example.splitlines()) - 1:
            r.add_break()

    add_heading(doc, "Snippet and content controls", 2)
    add_bullets(doc, [
        "Use noindex when a page should not appear in search. Ensure the crawler can access the page to see the directive; a robots block can prevent processing of noindex.",
        "Use nosnippet to prevent text or video snippets. Google states this also prevents the page content from being used as a direct input for AI Overviews and AI Mode.",
        "Use max-snippet to limit snippet length and data-nosnippet to exclude specific visible sections from snippets. Test the effect on traditional and AI visibility before broad use. [7]",
        "Do not place confidential, licensed, customer, checkout, order, or admin information on public pages and rely on snippet controls for security.",
    ])

    add_heading(doc, "Llms txt decision", 2)
    add_para(doc, "llms.txt is an open proposal that can provide a concise machine-readable site guide for systems that choose to consume it. Google explicitly says it does not use llms.txt for Search or generative Search visibility. Implement it only as a low-cost, documented experiment for known consumers or developer documentation use cases. It must not replace crawlable HTML, sitemaps, structured data, or platform-specific crawler controls. [1] [23]")

    page_break(doc)
    add_heading(doc, "Key Performance Indicators", 1)
    add_section_intro(doc, "Use a balanced set of leading, intermediate, and business measures. Every metric needs an owner, source, segment, frequency, and decision threshold.")

    add_table(
        doc,
        ["Layer", "Metric", "Definition and use"],
        [
            ["Eligibility", "Valid indexable share", "Canonical 200 URLs eligible for indexing divided by intended indexable URLs."],
            ["Eligibility", "Indexed share", "Indexed canonical URLs divided by submitted or intended URLs, segmented by template."],
            ["Eligibility", "Crawl error rate", "Error, blocked, challenge, or unexpected redirect responses divided by verified bot requests."],
            ["Experience", "Core Web Vitals pass rate", "URLs or origin traffic passing LCP, INP, and CLS at the 75th percentile."],
            ["Search", "Non-brand qualified impressions", "Impressions for approved non-brand query groups, excluding irrelevant demand."],
            ["Search", "Qualified organic visits", "Organic sessions reaching relevant landing pages, adjusted for bots and tracking quality."],
            ["Search", "Organic conversion and revenue", "Orders, leads, subscriptions, revenue, margin, and new-customer rate from organic visits and assisted paths."],
            ["Commerce", "Feed approval and parity", "Approved items plus percentage matching page price, stock, URL, shipping, and return data."],
            ["AEO", "Search activation rate", "Prompt runs that activate web search or expose sources divided by all runs."],
            ["AEO", "Brand mention rate", "Prompt runs mentioning the brand divided by all runs."],
            ["AEO", "Canonical citation rate", "Prompt runs citing at least one canonical owned URL divided by all runs."],
            ["AEO", "Answer accuracy", "Verified brand claims that are correct and current divided by reviewed claims."],
            ["AEO", "Citation diversity", "Distribution of citations across owned pages and third-party source types; use to detect dependence or concentration."],
            ["AEO", "AI referral quality", "Identifiable AI referrals by engaged session, conversion, assisted conversion, revenue, and new-customer rate."],
            ["Authority", "Relevant earned references", "New editorially independent, topically relevant links or mentions that accurately represent the brand."],
            ["Local", "Profile actions and local conversions", "Qualified calls, directions, website visits, messages, orders, and local landing conversions with platform limits noted."],
            ["Operations", "Issue closure time", "Median time to resolve P0 and P1 issues and verify the fix."],
            ["Operations", "Content freshness compliance", "Priority pages reviewed within their assigned interval with a recorded decision."],
        ],
        widths=[1.05, 1.75, 4.05],
        font_size=7.25,
    )

    add_heading(doc, "Reporting views", 2)
    add_bullets(doc, [
        "Executive view: qualified demand, revenue, brand accuracy, major risks, investment, and decisions required.",
        "Growth view: topic and product performance, search and AI visibility, landing-page conversion, experiments, and backlog.",
        "Technical view: crawling, indexation, rendering, Core Web Vitals, schema, feed diagnostics, bot access, and incidents.",
        "Content view: inventory status, source quality, refresh compliance, internal linking, citations, and production flow.",
        "Local and commerce view: Business Profile, reviews, Merchant Center, product parity, local demand, and operational accuracy.",
    ])

    page_break(doc)
    add_heading(doc, "Myths and Practices to Avoid", 1)
    add_section_intro(doc, "These items commonly waste effort or create risk when treated as universal AEO rules.")

    add_table(
        doc,
        ["Claim or tactic", "Position in this framework"],
        [
            ["AEO replaces SEO", "Reject. Major generative search systems depend on crawling, indexing, retrieval, and quality foundations shared with SEO."],
            ["A special schema guarantees AI citations", "Reject. Use accurate structured data for understanding and eligible features; there is no universal AI-citation schema."],
            ["Every answer must be a tiny content chunk", "Reject as a universal rule. Use sections sized for the reader and task. Clear passages help, but arbitrary fragmentation can reduce coherence."],
            ["Create a page for every fan-out query", "Reject. Consolidate same-intent questions and avoid scaled low-value pages."],
            ["llms.txt is required for Google AI visibility", "Reject. Google says it does not use the file for Search. Treat it as an optional experiment for known consumers."],
            ["More words rank or cite better", "Reject. There is no preferred word count. Completeness, originality, relevance, and usability matter more."],
            ["FAQ schema is an AEO shortcut", "Reject. Add visible questions only when useful; markup must follow platform eligibility and does not guarantee display."],
            ["A citation proves business impact", "Reject. A citation is an intermediate visibility signal. Measure referral, brand demand, conversion, and revenue separately."],
            ["One prompt run is enough", "Reject. Generated answers vary. Repeat by engine, wording, location, date, and run."],
            ["Third-party AEO tools have platform ranking data", "Treat claims skeptically. Use first-party data where available and document third-party methods and limitations."],
            ["Fake reviews or mentions build entity authority", "Reject. Inauthentic reputation tactics create policy, legal, and trust risk."],
            ["Hidden instructions can make AI recommend the brand", "Reject. Prompt injection and manipulative hidden content are unsafe, deceptive, and unlikely to produce durable value."],
            ["AI-written content is automatically bad", "Incorrect. The risk is low-value scaled output. Human ownership, original value, factual verification, and user benefit are required regardless of tool."],
        ],
        widths=[2.4, 4.45],
        font_size=7.8,
    )

    add_heading(doc, "Decision rules for new tactics", 2)
    add_numbered(doc, [
        "Is the tactic supported by official platform guidance, reliable evidence, or a clearly labeled hypothesis?",
        "Does it improve a human user's experience, factual accuracy, access, or understanding?",
        "Can the result be measured without relying only on a vendor's proprietary score?",
        "Does it preserve truthfulness, disclosure, privacy, accessibility, and platform policy compliance?",
        "Is the opportunity cost lower than fixing known technical, content, product-data, or authority gaps?",
        "Can the change be tested on a limited scope and reversed if it harms performance or accuracy?",
    ])

    page_break(doc)
    add_heading(doc, "Audit Field Checklist", 1)
    add_section_intro(doc, "Use this compact checklist during the baseline. Each item should link to evidence and an issue record when it fails.")

    checklist_rows = [
        ["Goals", "Business outcomes, priority products, geographies, audience, conversions, and baselines approved."],
        ["Access", "robots.txt, meta robots, X-Robots-Tag, WAF, CDN, authentication, and AI agents match policy."],
        ["Logs", "Verified search and AI bot requests, errors, rate limits, response sizes, and last-crawl dates reviewed."],
        ["Statuses", "Important URLs return correct 200, redirect, 404, or 410 responses without chains or soft 404s."],
        ["Canonicals", "Canonical URLs are indexable, equivalent, internally linked, in sitemaps, and selected by engines."],
        ["Sitemaps", "Only canonical indexable URLs, accurate lastmod, submitted and reconciled by template."],
        ["Rendering", "Main content, links, images, product facts, and structured data are present in rendered output."],
        ["Architecture", "Clear hierarchy, valuable categories, safe filters, short crawl paths, breadcrumbs, and no orphans."],
        ["Internal links", "Descriptive, crawlable, relevant, no broken links or redirecting destinations."],
        ["Metadata", "Unique accurate title and description, aligned main heading, canonical, open graph, and image metadata."],
        ["Experience", "Mobile usability, accessibility, semantic structure, forms, overlays, and Core Web Vitals reviewed."],
        ["Products", "Names, variants, images, price, INR, stock, preparation, shipping, returns, and policies complete."],
        ["Content", "Purpose, audience, originality, experience, expertise, accuracy, evidence, freshness, and next action recorded."],
        ["Answer readiness", "Direct answer, definitions, facts, comparisons, procedures, evidence, conditions, and limitations reviewed."],
        ["Entities", "Approved brand facts consistent across owned pages, schema, profiles, feeds, partners, and directories."],
        ["Schema", "Correct types, visible parity, required properties, stable IDs, validation, and monitoring."],
        ["Merchant", "Feed active, product parity, diagnostics, shipping, returns, identifiers, variants, and image quality."],
        ["Local", "Business Profile verified and current, reviews governed, NAP consistent, and local pages legitimate."],
        ["Authority", "Relevant links, mentions, profiles, partnerships, reviews, inaccuracies, and manipulation risks assessed."],
        ["AEO panel", "Prompts, engines, locations, repetitions, citations, accuracy, screenshots, and denominators documented."],
        ["Analytics", "Search, AI referrals, ecommerce, assisted conversions, revenue, logs, and annotations validated."],
        ["Governance", "Owners, release checks, refresh cadence, crawler policy, legal review, and incident path documented."],
    ]
    add_table(doc, ["Area", "Completion criterion"], checklist_rows, widths=[1.2, 5.65], font_size=7.75)

    add_heading(doc, "Recommended first backlog for Bean Leaf Brew", 2)
    add_para(doc, "This list is a planning hypothesis based on the current ecommerce structure, not the result of the full audit. The audit should confirm scope, severity, and order.")
    add_table(
        doc,
        ["Priority", "Work item", "Acceptance signal"],
        [
            ["P0", "Confirm public versus private route indexation and crawler access, including admin, API, checkout, success, and product pages.", "Approved rules; no accidental public indexing; intended product and content pages accessible."],
            ["P1", "Create XML sitemap, robots policy, canonical rules, and search console verification where missing or incomplete.", "Valid submitted canonical URLs with clean inspection results."],
            ["P1", "Implement Organization, WebSite, BreadcrumbList, Product, and Offer data from authoritative fields.", "Rendered, valid markup that matches visible content and feed data."],
            ["P1", "Set up or reconcile Merchant Center, shipping, returns, product images, price, stock, and variants.", "Approved products with low diagnostics and consistent landing pages."],
            ["P1", "Complete public shipping, cancellation, refund, return, privacy, and contact information.", "Policies are current, easy to find, and linked from relevant purchase paths."],
            ["P1", "Improve unique metadata, canonical URLs, image metadata, and product detail completeness.", "Priority product pages pass template and content acceptance criteria."],
            ["P2", "Create useful category landing pages and stronger internal linking between products and guides.", "Strategic products reachable through clear crawlable paths with distinct category intent."],
            ["P2", "Launch brewing, selection, comparison, storage, office premix, and local delivery answer assets.", "Original, verified content published with product paths and baseline performance tracking."],
            ["P2", "Align Business Profile and local facts; establish review and profile update procedures.", "Consistent public information and a documented monthly review."],
            ["P2", "Create the prompt panel and monthly AI visibility and brand accuracy review.", "Repeatable multi-engine benchmark with citations, accuracy, and referral reporting."],
            ["P3", "Test llms.txt only if a defined consumer or low-cost hypothesis exists.", "Experiment documented separately from core search KPIs."],
        ],
        widths=[0.7, 4.15, 2.0],
        font_size=7.55,
    )

    page_break(doc)
    add_heading(doc, "Research Sources", 1)
    add_section_intro(doc, "The framework prioritizes first-party platform documentation and labels emerging research or proposals separately. Sources were accessed 30 October 2026.")

    sources = [
        ("1", "Google Search Central. Optimizing your website for generative AI features on Google Search. Updated 10 July 2026.", "https://developers.google.com/search/docs/fundamentals/ai-optimization-guide"),
        ("2", "Google Search Central. Google Search Essentials.", "https://developers.google.com/search/docs/essentials"),
        ("3", "Google Search Central. Search Engine Optimization Starter Guide.", "https://developers.google.com/search/docs/fundamentals/seo-starter-guide"),
        ("4", "Google Search Central. Creating helpful reliable people first content.", "https://developers.google.com/search/docs/fundamentals/creating-helpful-content"),
        ("5", "Google Search Central. Spam policies for Google web search.", "https://developers.google.com/search/docs/essentials/spam-policies"),
        ("6", "Google Search Central. Overview of crawling and indexing topics.", "https://developers.google.com/search/docs/crawling-indexing"),
        ("7", "Google Search Central. Robots meta tags and snippet controls.", "https://developers.google.com/search/docs/crawling-indexing/robots-meta-tag"),
        ("8", "Google Search Central. What is URL canonicalization.", "https://developers.google.com/search/docs/crawling-indexing/canonicalization"),
        ("9", "Google Search Central. General structured data guidelines.", "https://developers.google.com/search/docs/appearance/structured-data/sd-policies"),
        ("10", "Google Search Central. Introduction to Product structured data and merchant listings.", "https://developers.google.com/search/docs/appearance/structured-data/product"),
        ("11", "Google Search Central. Organization structured data.", "https://developers.google.com/search/docs/appearance/structured-data/organization"),
        ("12", "Google Search Central. LocalBusiness structured data.", "https://developers.google.com/search/docs/appearance/structured-data/local-business"),
        ("13", "Google Search Central. SEO best practices for ecommerce sites.", "https://developers.google.com/search/docs/specialty/ecommerce"),
        ("14", "web.dev. Web Vitals and current Core Web Vitals thresholds.", "https://web.dev/articles/vitals"),
        ("15", "Google Search Central. Generative AI performance reports in Search Console. 3 June 2026.", "https://developers.google.com/search/blog/2026/06/gen-ai-performance-reports"),
        ("16", "Google Search Central. How to use Search Console.", "https://developers.google.com/search/docs/monitor-debug/search-console-start"),
        ("17", "Bing Webmaster Tools. Webmaster Guidelines.", "https://www.bing.com/webmasters/help/webmaster-guidelines-30fba23a"),
        ("18", "Bing Webmaster Blog. Introducing AI Performance in Bing Webmaster Tools. 10 February 2026.", "https://blogs.bing.com/webmaster/2026/2/Introducing-AI-Performance-in-Bing-Webmaster-Tools-Public-Preview/"),
        ("19", "OpenAI Developers. Overview of OpenAI crawlers.", "https://developers.openai.com/api/docs/bots"),
        ("20", "Perplexity Documentation. Perplexity Crawlers.", "https://docs.perplexity.ai/docs/resources/perplexity-crawlers"),
        ("21", "Claude Help Center. Anthropic crawlers and website owner controls. 7 April 2026.", "https://support.claude.com/en/articles/8896518-does-anthropic-crawl-data-from-the-web-and-how-can-site-owners-block-the-crawler"),
        ("22", "Schema.org. Schemas and vocabulary documentation.", "https://schema.org/docs/schemas.html"),
        ("23", "llms.txt. The llms.txt proposal version 2. Modified 10 August 2026.", "https://llmstxt.org/"),
        ("24", "Martinez O. Optimizing Visibility in Generative Engines. Critical survey of GEO research from 2023 to 2026. Preprint.", "https://arxiv.org/abs/2607.14035"),
        ("25", "Google Search Central. Guidance on third party SEO tools services and advice. 2026.", "https://developers.google.com/search/docs/fundamentals/third-party-seo"),
        ("26", "Google Business Profile Help. Tips to improve local ranking on Google.", "https://support.google.com/business/answer/7091"),
        ("27", "Google Merchant Center Help. Free listings for products.", "https://support.google.com/merchants/answer/13889434"),
        ("28", "Aggarwal P and others. GEO Generative Engine Optimization. Foundational research paper and GEO benchmark. 2023.", "https://arxiv.org/abs/2311.09735"),
    ]
    for num, label, url in sources:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(f"{num}. {label} ")
        set_run_font(r, size=8.6)
        add_hyperlink(p, "Source", url)

    add_heading(doc, "Research interpretation", 2)
    add_para(doc, "The original GEO paper reported visibility gains in a controlled benchmark, but later reviews caution that results depend on an already retrieved context and do not prove durable organic discovery, traffic, or revenue. This framework therefore uses research-derived formatting and measurement ideas as testable practices, not guarantees. [24] [28]")

    add_heading(doc, "Final planning note", 2)
    add_para(doc, "This handbook defines the complete practice set. The next step is to apply it to the live site and connected platform accounts, capture evidence, calculate the score, and convert the findings into an owned implementation backlog. The audit should preserve a distinction between verified facts, platform guidance, analysis, and hypotheses.")

    configure_sections(doc)
    return doc


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.core_properties.title = "AEO and SEO Audit and Implementation Framework"
    doc.core_properties.subject = "Comprehensive SEO and answer engine optimization audit and implementation practices for Bean Leaf Brew"
    doc.core_properties.author = "Bean Leaf Brew"
    doc.core_properties.keywords = "AEO, SEO, GEO, AI search, ecommerce, local SEO, structured data, audit"
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
